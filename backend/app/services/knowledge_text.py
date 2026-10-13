"""Pipeline de texto da base de conhecimento (§29):

Upload → Extração → Normalização → Chunking (→ Embedding → Vector DB em `ai/rag.py`).
"""

import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

#: Formatos aceitos no upload de artigos (§29).
KNOWLEDGE_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

#: Tamanho alvo de cada trecho indexado (caracteres) e sobreposição entre trechos.
CHUNK_SIZE = 900
CHUNK_OVERLAP = 150


class ExtractionError(Exception):
    def __init__(self, message: str, code: str = "EXTRACTION_FAILED") -> None:
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass(frozen=True)
class ExtractedDocument:
    title: str
    text: str


def extract_text(file_name: str, data: bytes) -> ExtractedDocument:
    """Extrai texto puro de PDF, DOCX, TXT ou MD. O título vem do nome do arquivo."""
    extension = PurePath(file_name).suffix.lower()
    title = PurePath(file_name).stem.replace("_", " ").replace("-", " ").strip() or "Documento"

    if extension in (".txt", ".md"):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ExtractionError(
                "O arquivo de texto deve estar em UTF-8.", "INVALID_ENCODING"
            ) from exc
        if extension == ".md":
            heading = re.search(r"^#\s+(.+)$", text, flags=re.MULTILINE)
            if heading:
                title = heading.group(1).strip()
    elif extension == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractionError("O arquivo não é um PDF válido.", "FILE_CONTENT_MISMATCH")
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except (PdfReadError, ValueError) as exc:
            raise ExtractionError("Não foi possível ler o PDF.") from exc
    elif extension == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ExtractionError("O arquivo não é um DOCX válido.", "FILE_CONTENT_MISMATCH")
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:  # python-docx levanta vários tipos para arquivos corrompidos
            raise ExtractionError("Não foi possível ler o DOCX.") from exc
        text = "\n\n".join(paragraph.text for paragraph in document.paragraphs)
    else:
        raise ExtractionError(
            "Formato não suportado. Envie PDF, DOCX, TXT ou MD.", "FILE_TYPE_NOT_ALLOWED"
        )

    text = normalize_text(text)
    if len(text) < 20:
        raise ExtractionError(
            "Não foi encontrado texto suficiente no arquivo "
            "(PDFs digitalizados não são suportados).",
            "EMPTY_DOCUMENT",
        )
    return ExtractedDocument(title=title[:200], text=text)


def normalize_text(text: str) -> str:
    """Unicode NFC, sem caracteres de controle, espaços e linhas em branco consolidados."""
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(char for char in text if char in "\n\t" or unicodedata.category(char)[0] != "C")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, *, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Divide em trechos de ~`size` caracteres respeitando parágrafos e frases.

    Parágrafos são agrupados até o limite; parágrafos longos são quebrados em
    frases. Trechos consecutivos repetem o final do anterior (`overlap`) para
    que uma resposta que cruza a fronteira não se perca.
    """
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    pieces: list[str] = []
    for paragraph in paragraphs:
        if len(paragraph) <= size:
            pieces.append(paragraph)
            continue
        sentences = re.split(r"(?<=[.!?])\s+", paragraph)
        current = ""
        for sentence in sentences:
            while len(sentence) > size:  # frase gigante: corte duro
                pieces.append(sentence[:size])
                sentence = sentence[size:]
            if current and len(current) + len(sentence) + 1 > size:
                pieces.append(current)
                current = sentence
            else:
                current = f"{current} {sentence}".strip()
        if current:
            pieces.append(current)

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + len(piece) + 2 > size:
            chunks.append(current)
            tail = current[-overlap:]
            # Começa a sobreposição numa fronteira de palavra.
            tail = tail[tail.find(" ") + 1 :] if " " in tail else tail
            current = f"{tail}\n\n{piece}" if overlap else piece
        else:
            current = f"{current}\n\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks
