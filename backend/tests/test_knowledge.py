"""Base de conhecimento (§29) e RAG (§28): extração, chunking, busca e permissões."""

import io
from collections.abc import Callable
from typing import Any

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Company, KnowledgeArticle, KnowledgeChunk, User
from app.models.enums import ArticleStatus, RoleName
from app.services.ai.embeddings import HashingEmbedder, cosine, tokenize
from app.services.ai.rag import RAGService
from app.services.ai.vector_store import (
    InMemoryVectorStore,
    VectorRecord,
    VectorStoreError,
    collection_name,
)
from app.services.knowledge_text import ExtractionError, chunk_text, extract_text

VPN_ARTICLE = {
    "title": "Como resolver falha de autenticação na VPN",
    "summary": "Passo a passo quando a VPN recusa as credenciais.",
    "content": (
        "Quando a VPN mostrar 'falha de autenticação', verifique se a sua senha de rede "
        "não expirou. Em seguida atualize o cliente VPN para a versão mais recente e "
        "reinicie o computador. Se o erro continuar, abra um chamado na categoria VPN."
    ),
    "status": "PUBLISHED",
}
PRINTER_ARTICLE = {
    "title": "Impressora atolando papel",
    "content": (
        "Abra a tampa traseira da impressora, retire o papel preso com cuidado e verifique "
        "se o rolete de alimentação está limpo. Use somente papel A4 de 75 g/m²."
    ),
    "status": "PUBLISHED",
}


def _error(response: Any) -> dict[str, Any]:
    return response.json()["error"]


@pytest.fixture
def company(make_company: Callable[..., Company]) -> Company:
    return make_company(name="Empresa Alfa")


@pytest.fixture
def people(
    company: Company, make_user: Callable[..., User], db_session: Session
) -> dict[str, User]:
    users = {
        "ana": make_user(company, role=RoleName.USER),
        "teo": make_user(company, role=RoleName.TECHNICIAN),
        "alice": make_user(company, role=RoleName.ADMIN),
    }
    db_session.commit()
    return users


def create(client: TestClient, headers: dict[str, str], payload: dict[str, Any]) -> dict[str, Any]:
    response = client.post("/api/v1/knowledge", json=payload, headers=headers)
    assert response.status_code == 201, response.text
    return response.json()


class TestTextPipeline:
    def test_extracts_txt_and_md_with_title(self) -> None:
        assert (
            extract_text("guia_vpn.txt", b"Passo a passo da VPN corporativa.").title == "guia vpn"
        )
        md = extract_text("x.md", b"# Configurar o e-mail\n\nAbra o Outlook e ...")
        assert md.title == "Configurar o e-mail"

    def test_extracts_docx(self) -> None:
        document = Document()
        document.add_paragraph("Como mapear a unidade de rede do financeiro.")
        document.add_paragraph("Abra o Explorador de Arquivos e clique em Mapear unidade.")
        buffer = io.BytesIO()
        document.save(buffer)

        result = extract_text("rede.docx", buffer.getvalue())

        assert "Mapear unidade" in result.text

    @pytest.mark.parametrize(
        ("name", "data", "code"),
        [
            ("falso.pdf", b"isto nao e pdf", "FILE_CONTENT_MISMATCH"),
            ("falso.docx", b"isto nao e docx", "FILE_CONTENT_MISMATCH"),
            ("planilha.xlsx", b"PK\x03\x04", "FILE_TYPE_NOT_ALLOWED"),
            ("vazio.txt", b"   ", "EMPTY_DOCUMENT"),
            ("latin1.txt", "ação".encode("latin-1") * 10, "INVALID_ENCODING"),
        ],
    )
    def test_rejects_invalid_files(self, name: str, data: bytes, code: str) -> None:
        with pytest.raises(ExtractionError) as error:
            extract_text(name, data)
        assert error.value.code == code

    def test_chunks_respect_the_size_and_overlap(self) -> None:
        text = "\n\n".join(f"Parágrafo {i}. " + "texto " * 60 for i in range(10))

        chunks = chunk_text(text, size=500, overlap=80)

        assert len(chunks) > 3
        assert all(len(chunk) <= 500 + 80 + 2 for chunk in chunks)
        # Sobreposição: cada trecho começa com o final do anterior.
        assert chunks[1].split("\n\n")[0] in chunks[0]


def test_embeddings_bring_related_texts_closer() -> None:
    embedder = HashingEmbedder()
    question, vpn, printer = embedder.embed(
        [
            "minha vpn não conecta, diz falha de autenticação",
            VPN_ARTICLE["content"],
            PRINTER_ARTICLE["content"],
        ]
    )

    assert cosine(question, vpn) > cosine(question, printer) + 0.1


@pytest.mark.parametrize(
    ("typed", "written"),
    [
        ("emails", "e-mail"),
        ("wifi", "Wi-Fi"),
        ("pendrive", "pen-drive"),
        ("PC", "computador"),
        ("funcionando", "funcionar"),
        ("impressão", "impressora"),
        ("configurações", "configurar"),
    ],
)
def test_tokenizer_matches_what_people_type_with_how_articles_are_written(
    typed: str, written: str
) -> None:
    assert tokenize(typed) == tokenize(written)


class TestRAG:
    def test_search_returns_only_published_articles_of_the_company(
        self,
        client: TestClient,
        db_session: Session,
        people: dict[str, User],
        make_company: Callable[..., Company],
        make_user: Callable[..., User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        admin = auth_headers(people["alice"])
        vpn = create(client, admin, VPN_ARTICLE)
        create(client, admin, PRINTER_ARTICLE)
        create(client, admin, {**VPN_ARTICLE, "title": "Rascunho sobre VPN", "status": "DRAFT"})
        foreign_admin = make_user(make_company(name="Beta"), role=RoleName.ADMIN)
        db_session.commit()
        create(client, auth_headers(foreign_admin), {**VPN_ARTICLE, "title": "VPN da Beta"})

        response = client.get(
            "/api/v1/knowledge/suggest?q=vpn falha de autenticação",
            headers=auth_headers(people["ana"]),
        )

        assert response.status_code == 200
        results = response.json()
        assert [r["article_id"] for r in results][0] == vpn["id"]
        titles = {r["title"] for r in results}
        assert "Rascunho sobre VPN" not in titles
        assert "VPN da Beta" not in titles

    def test_falls_back_to_database_scan_when_the_vector_store_fails(
        self, db_session: Session, people: dict[str, User], company: Company
    ) -> None:
        class BrokenStore(InMemoryVectorStore):
            def query(self, company_id: int, embedding: list[float], k: int) -> list[Any]:
                raise VectorStoreError("fora do ar")

        rag = RAGService(db_session, BrokenStore(), HashingEmbedder())
        article = KnowledgeArticle(
            company_id=company.id,
            title=VPN_ARTICLE["title"],
            content=VPN_ARTICLE["content"],
            status=ArticleStatus.PUBLISHED,
        )
        db_session.add(article)
        db_session.flush()
        rag.index_article(article)

        results = rag.search(company.id, "vpn com falha de autenticação")

        assert results and results[0].article_id == article.id

    def test_rebuild_replaces_stale_vectors(
        self, db_session: Session, people: dict[str, User], company: Company
    ) -> None:
        store = InMemoryVectorStore()
        rag = RAGService(db_session, store, HashingEmbedder())
        article = KnowledgeArticle(
            company_id=company.id,
            title=VPN_ARTICLE["title"],
            content=VPN_ARTICLE["content"],
            status=ArticleStatus.PUBLISHED,
        )
        db_session.add(article)
        db_session.flush()
        rag.index_article(article)
        # Vetor de uma versão antiga do embedder, sem artigo correspondente.
        store.upsert(company.id, [VectorRecord("old", [0.0] * 8, "antigo", 999, 0)])

        chunks = rag.rebuild_company(company.id)

        records = store._data[collection_name(company.id)]
        assert chunks == len(records) >= 1
        assert {record.article_id for record in records.values()} == {article.id}
        assert rag.search(company.id, "vpn com falha de autenticação")[0].article_id == article.id

    def test_rejects_queries_without_company(self, db_session: Session) -> None:
        rag = RAGService(db_session, InMemoryVectorStore(), HashingEmbedder())
        with pytest.raises(ValueError):
            rag.search(0, "vpn")


class TestAPI:
    def test_admin_publishes_and_the_article_is_indexed(
        self,
        client: TestClient,
        db_session: Session,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        article = create(client, auth_headers(people["alice"]), VPN_ARTICLE)

        assert article["indexed_at"] is not None
        assert article["author"]["id"] == people["alice"].id
        chunks = db_session.scalars(
            select(KnowledgeChunk).where(KnowledgeChunk.article_id == article["id"])
        ).all()
        assert chunks and all(chunk.company_id == people["alice"].company_id for chunk in chunks)

    def test_unpublishing_removes_from_the_index(
        self,
        client: TestClient,
        db_session: Session,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        admin = auth_headers(people["alice"])
        article = create(client, admin, VPN_ARTICLE)

        updated = client.put(
            f"/api/v1/knowledge/{article['id']}", json={"status": "ARCHIVED"}, headers=admin
        ).json()

        assert updated["indexed_at"] is None
        assert not db_session.scalars(
            select(KnowledgeChunk).where(KnowledgeChunk.article_id == article["id"])
        ).all()
        suggest = client.get("/api/v1/knowledge/suggest?q=vpn autenticação", headers=admin)
        assert suggest.json() == []

    @pytest.mark.parametrize("role", ["ana", "teo"])
    def test_only_admin_manages_articles(
        self,
        client: TestClient,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
        role: str,
    ) -> None:
        response = client.post(
            "/api/v1/knowledge", json=VPN_ARTICLE, headers=auth_headers(people[role])
        )

        assert response.status_code == 403
        assert _error(response)["code"] == "KNOWLEDGE_FORBIDDEN"

    def test_users_see_only_published_articles(
        self,
        client: TestClient,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        admin = auth_headers(people["alice"])
        create(client, admin, VPN_ARTICLE)
        draft = create(client, admin, {**PRINTER_ARTICLE, "status": "DRAFT"})

        as_user = client.get("/api/v1/knowledge", headers=auth_headers(people["ana"])).json()
        as_staff = client.get("/api/v1/knowledge", headers=auth_headers(people["teo"])).json()

        assert as_user["total"] == 1
        assert as_staff["total"] == 2
        hidden = client.get(f"/api/v1/knowledge/{draft['id']}", headers=auth_headers(people["ana"]))
        assert hidden.status_code == 404

    def test_search_category_and_sorting(
        self,
        client: TestClient,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        admin = auth_headers(people["alice"])
        create(client, admin, VPN_ARTICLE)
        printer = create(client, admin, {**PRINTER_ARTICLE, "is_featured": True})

        found = client.get("/api/v1/knowledge?search=rolete", headers=admin).json()
        popular = client.get("/api/v1/knowledge?sort=popular", headers=admin).json()

        assert [item["title"] for item in found["items"]] == ["Impressora atolando papel"]
        assert popular["items"][0]["id"] == printer["id"]  # destacado primeiro

    def test_views_and_feedback_are_counted(
        self,
        client: TestClient,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        article = create(client, auth_headers(people["alice"]), VPN_ARTICLE)
        user = auth_headers(people["ana"])

        client.get(f"/api/v1/knowledge/{article['id']}", headers=user)
        viewed = client.get(f"/api/v1/knowledge/{article['id']}", headers=user).json()
        rated = client.post(
            f"/api/v1/knowledge/{article['id']}/feedback", json={"helpful": False}, headers=user
        ).json()

        assert viewed["view_count"] == 2
        assert (rated["helpful_count"], rated["not_helpful_count"]) == (0, 1)

    def test_upload_creates_a_draft_for_review(
        self,
        client: TestClient,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        document = Document()
        document.add_paragraph("Para acessar o Wi-Fi corporativo use a rede EMPRESA-5G.")
        buffer = io.BytesIO()
        document.save(buffer)

        response = client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": ("wifi_corporativo.docx", buffer.getvalue(), "application/octet-stream")
            },
            headers=auth_headers(people["alice"]),
        )

        assert response.status_code == 201
        body = response.json()
        assert (body["status"], body["source_type"], body["title"]) == (
            "DRAFT",
            "UPLOAD",
            "wifi corporativo",
        )
        assert "EMPRESA-5G" in body["content"]
        assert body["indexed_at"] is None  # rascunho não entra no RAG

    def test_upload_rejects_invalid_files(
        self,
        client: TestClient,
        people: dict[str, User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        response = client.post(
            "/api/v1/knowledge/upload",
            files={"file": ("virus.exe", b"MZ\x90", "application/octet-stream")},
            headers=auth_headers(people["alice"]),
        )

        assert response.status_code == 400
        assert _error(response)["code"] == "FILE_TYPE_NOT_ALLOWED"

    def test_articles_of_another_company_are_not_found(
        self,
        app: FastAPI,
        client: TestClient,
        db_session: Session,
        people: dict[str, User],
        make_company: Callable[..., Company],
        make_user: Callable[..., User],
        auth_headers: Callable[[User], dict[str, str]],
    ) -> None:
        article = create(client, auth_headers(people["alice"]), VPN_ARTICLE)
        outsider = make_user(make_company(name="Beta"), role=RoleName.ADMIN)
        db_session.commit()
        headers = auth_headers(outsider)

        assert client.get(f"/api/v1/knowledge/{article['id']}", headers=headers).status_code == 404
        assert (
            client.delete(f"/api/v1/knowledge/{article['id']}", headers=headers).status_code == 404
        )


def test_chroma_circuit_breaker_fails_fast_after_an_error() -> None:
    """Com o Chroma fora do ar, só a primeira chamada paga o timeout (regressão)."""
    import time

    from app.services.ai.vector_store import ChromaVectorStore

    store = ChromaVectorStore("http://127.0.0.1:9")  # porta sem serviço
    with pytest.raises(VectorStoreError):
        store.query(1, [0.0] * 512, 3)

    started = time.perf_counter()
    for _ in range(5):
        with pytest.raises(VectorStoreError):
            store.query(1, [0.0] * 512, 3)
    assert time.perf_counter() - started < 0.5
